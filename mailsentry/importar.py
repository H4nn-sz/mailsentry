"""Importación masiva de personal desde Excel (.xlsx), CSV o texto copiado de una hoja de cálculo."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from .util import normalizar

COLUMNAS = {  # nombre de columna normalizado -> campo
    "nombre": "nombre", "nombres": "nombre", "nombre y apellido": "nombre", "nombre completo": "nombre",
    "apellidos y nombres": "nombre", "colaborador": "nombre", "empleado": "nombre",
    "correo": "email", "email": "email", "e-mail": "email", "mail": "email", "correo electronico": "email",
    "correo de la empresa": "email", "correo corporativo": "email",
    "cargo": "cargo", "puesto": "cargo", "posicion": "cargo",
    "area": "departamento", "departamento": "departamento", "gerencia": "departamento", "unidad": "departamento",
    "directivo": "es_vip", "vip": "es_vip", "es directivo": "es_vip",
}
SI = {"si", "s", "x", "1", "true", "verdadero", "yes", "y"}
RE_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PLANTILLA = ("nombre,correo,cargo,area,directivo\n"
             "Ana Torres,ana.torres@miempresa.com.pe,Gerente General,Gerencia,si\n"
             "Luis Pérez,luis.perez@miempresa.com.pe,Contador,Finanzas,no\n")


@dataclass
class Persona:
    email: str
    nombre: str = ""
    cargo: str = ""
    departamento: str = ""
    es_vip: bool = False


@dataclass
class ResultadoImportacion:
    personas: list[Persona] = field(default_factory=list)
    omitidas: list[str] = field(default_factory=list)  # motivo por fila


def _filas_xlsx(datos: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    return [["" if v is None else str(v) for v in fila] for fila in hoja.iter_rows(values_only=True)]


def _filas_texto(texto: str) -> list[list[str]]:
    texto = texto.lstrip("﻿").strip()
    if not texto:
        return []
    muestra = texto[:2000]
    separador = "\t" if "\t" in muestra else (";" if muestra.count(";") > muestra.count(",") else ",")
    return [fila for fila in csv.reader(io.StringIO(texto), delimiter=separador)]


def leer(nombre_archivo: str = "", datos: bytes = b"", texto: str = "") -> ResultadoImportacion:
    """Acepta un .xlsx, un .csv o texto pegado desde Excel. La primera fila puede ser de títulos o no."""
    if nombre_archivo.lower().endswith((".xlsx", ".xlsm")):
        filas = _filas_xlsx(datos)
    elif datos:
        try:
            contenido = datos.decode("utf-8-sig")
        except UnicodeDecodeError:
            contenido = datos.decode("latin-1")  # CSV guardado por Excel en Windows
        filas = _filas_texto(contenido)
    else:
        filas = _filas_texto(texto)
    filas = [f for f in filas if any(c.strip() for c in f)]
    resultado = ResultadoImportacion()
    if not filas:
        return resultado

    # ¿La primera fila son títulos? Si reconoce columnas, las usa; si no, asume: nombre, correo, cargo, área, directivo
    # Títulos como "¿Directivo?", "Correo*" o "Área:" se limpian antes de reconocerlos
    titulos = [COLUMNAS.get(re.sub(r"[^a-z0-9 -]", "", normalizar(c)).strip()) for c in filas[0]]
    if "email" in titulos:
        mapa = titulos
        filas = filas[1:]
    else:
        mapa = ["nombre", "email", "cargo", "departamento", "es_vip"]
        if not any(RE_EMAIL.match(c.strip()) for c in filas[0]):  # títulos no reconocidos
            filas = filas[1:]

    vistos = set()
    for numero, fila in enumerate(filas, start=2):
        valores = {campo: (fila[i].strip() if i < len(fila) else "") for i, campo in enumerate(mapa) if campo}
        email = valores.get("email", "").lower().strip("<> ")
        if not RE_EMAIL.match(email):
            resultado.omitidas.append(f"fila {numero}: correo inválido «{valores.get('email', '')}»")
            continue
        if email in vistos:
            resultado.omitidas.append(f"fila {numero}: {email} repetido")
            continue
        vistos.add(email)
        resultado.personas.append(Persona(
            email=email, nombre=valores.get("nombre", "")[:200], cargo=valores.get("cargo", "")[:120],
            departamento=valores.get("departamento", "")[:120],
            es_vip=normalizar(valores.get("es_vip", "")) in SI,
        ))
    return resultado
